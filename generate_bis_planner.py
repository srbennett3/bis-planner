#!/usr/bin/env python3
"""
BIS Planner — TBC Classic Gear Spreadsheet Generator
Created by Steven Bennett 2026

Two-step workflow:
  1. python3 generate_bis_planner.py --build-db         # fetch item stats from Wowhead (re-fetches stale cache)
  2. python3 generate_bis_planner.py paladin             # generate CSV + Excel (no API calls)

Tooltip parsing strips (N) Set : … spell links so conditional set bonuses are not stored as base stats (fixes gear score / % Upgrade vs Pawn).

Loon (BIS) guides: AddonReference/Loon/*.lua. Pawn weights: AddonReference/Pawn/ClassicHawsJon.lua → data/pawn_scales_tbc.json.
TBC gem lists: Pawn/GemsBurningCrusade.lua. Excel adds Current Equipment ideal-gem columns (Apps Script), Weights row (Pawn snapshot + custom stash), Gear Score, GemDB + ItemDB socket columns; BIS Planner Stat Comparison + % Upgrade.

Supported classes: druid, hunter, mage, paladin, priest, rogue, shaman, warlock, warrior

Requires: openpyxl  (pip install openpyxl)
"""

import argparse
import concurrent.futures
import csv
from collections import Counter
import json
import os
import re
import ssl
import sys
import threading
import time
import urllib.error
import urllib.request

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ADDON_REFERENCE_DIR = os.path.join(SCRIPT_DIR, "AddonReference")
LOON_GUIDES_DIR = os.path.join(ADDON_REFERENCE_DIR, "Loon")
PAWN_ADDON_DIR = os.path.join(ADDON_REFERENCE_DIR, "Pawn")
PAWN_CLASSIC_HAWS_PATH = os.path.join(PAWN_ADDON_DIR, "ClassicHawsJon.lua")
DATA_DIR = os.path.join(SCRIPT_DIR, "data")
PAWN_SCALES_JSON_PATH = os.path.join(DATA_DIR, "pawn_scales_tbc.json")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "output")
ITEM_DB_PATH = os.path.join(DATA_DIR, "item_database.json")
GEAR_ORDER = [
    "Head", "Shoulder", "Back", "Chest", "Wrist", "Hands",
    "Waist", "Legs", "Feet", "Neck", "Ring", "Trinket",
    "Main Hand", "Off Hand", "Two Hand", "Ranged/Relic",
]
# Current Equipment rows (two ring + two trinket slots); planner CSV still uses "Ring" / "Trinket".
CE_GEAR_ORDER = [
    "Head", "Shoulder", "Back", "Chest", "Wrist", "Hands",
    "Waist", "Legs", "Feet", "Neck",
    "Ring 1", "Ring 2", "Trinket 1", "Trinket 2",
    "Main Hand", "Off Hand", "Two Hand", "Ranged/Relic",
]
CE_ROW_TO_ITEMDB_SLOT = {
    "Ring 1": "Ring",
    "Ring 2": "Ring",
    "Trinket 1": "Trinket",
    "Trinket 2": "Trinket",
}
ACQ_ORDER = [
    "Auction House", "Dungeon Drop", "Dungeon Token",
    "Quest", "Quest (Dung)", "Reputation", "PvP",
    "World Drop", "Raid Drop", "Vendor", "Other",
]

DUNGEON_KEYWORDS = [
    "Shattered Halls", "Steamvault", "Botanica", "Shadow Labyrinth",
    "Auchenai Crypts", "Black Morass", "Slave Pens", "Mechanar",
    "Arcatraz", "Hellfire Ramparts", "Blood Furnace", "Underbog",
    "Mana-Tombs", "Mana Tombs", "Sethekk Halls", "Old Hillsbrad",
    "Caverns of Time", "Tempest Keep",
]

LOCATION_ABBREVS = {
    "The Shattered Halls": "SH", "Shattered Halls": "SH",
    "The Steamvault": "SV", "Steamvault": "SV",
    "The Botanica": "Bot", "Botanica": "Bot",
    "Shadow Labyrinth": "SL", "Auchenai Crypts": "AC",
    "The Black Morass": "BM", "Black Morass": "BM",
    "The Slave Pens": "SP", "Slave Pens": "SP",
    "The Mechanar": "Mech", "Mechanar": "Mech",
    "The Arcatraz": "Arc", "Arcatraz": "Arc",
    "Hellfire Ramparts": "Ramp",
    "The Blood Furnace": "BF", "Blood Furnace": "BF",
    "The Underbog": "UB", "Underbog": "UB",
    "Mana-Tombs": "MT", "Mana Tombs": "MT",
    "Sethekk Halls": "Seth",
    "Old Hillsbrad Foothills": "OHF",
    "Opening of the Dark Portal": "BM",
    "Caverns of Time": "CoT", "Tempest Keep": "TK",
    "Netherstorm": "NS", "Blade's Edge Mountains": "BEM",
    "Shadowmoon Valley": "SMV", "Nagrand": "Nag",
    "Zangarmarsh": "ZM", "Terokkar Forest": "TF",
    "Hellfire Peninsula": "HFP", "Auchindoun": "Auch",
    "Blackrock Spire": "BRS", "Naxxramas": "Naxx",
    "Ahn'Qiraj": "AQ", "Temple of Ahn'Qiraj": "AQ40",
    "Zul'Gurub": "ZG",
}

WOWHEAD_DATA_ENV = 5
FETCH_DELAY = 0.15
FETCH_WORKERS = 4

STAT_COLUMNS = [
    "Armor", "Str", "Agi", "Sta", "Int", "Spi",
    "Healing", "Spell Dmg", "AP", "MP5",
    "Defense", "Dodge", "Parry", "Block Rating", "Block Value",
    "Hit", "Crit", "Spell Hit", "Spell Crit", "Haste",
    "Resilience",  # high usage in TBC PvP gear; maps to Pawn ResilienceRating
]

# ItemDB: Name, Slot, then socket counts, then STAT_COLUMNS, Attributes.
# Must stay in sync with apps_script.gs (comment CE ↔ ITEMDB).
ITEMDB_SOCK_COL_FIRST = 3  # 1-based: SockR
NUM_ITEMDB_SOCKET_COLS = 4
ITEMDB_FIRST_STAT_COL = ITEMDB_SOCK_COL_FIRST + NUM_ITEMDB_SOCKET_COLS  # 7

# Current Equipment: A=Gear, B=Item, ideal gem summary (Apps Script), stats, Gear Score; then hidden stash.
CE_GEM_HEADERS = [
    "Ideal Gem",
    "Ideal Meta",
    "Total Gem Stats",
]
NUM_CE_GEM_COLS = len(CE_GEM_HEADERS)
CE_WEIGHT_ROW_LABEL = "Weights"
CE_WEIGHT_MODE_PAWN = "Pawn Default"
CE_WEIGHT_MODE_CUSTOM = "Custom"
# 1-based column indices (must match generate_bis_planner.py export_xlsx & apps_script.gs).
CE_GEM_FIRST_COL = 3
CE_GEM_LAST_COL = CE_GEM_FIRST_COL + NUM_CE_GEM_COLS - 1
CE_STAT_FIRST_COL = CE_GEM_LAST_COL + 1
CE_STAT_LAST_COL = CE_STAT_FIRST_COL + len(STAT_COLUMNS) - 1
CE_GEAR_SCORE_COL = CE_STAT_LAST_COL + 1
CE_PAWN_SNAPSHOT_FIRST_COL = CE_GEAR_SCORE_COL + 1
CE_CUSTOM_STASH_FIRST_COL = CE_PAWN_SNAPSHOT_FIRST_COL + len(STAT_COLUMNS)
CE_META_WEIGHT_COL = CE_CUSTOM_STASH_FIRST_COL + len(STAT_COLUMNS)

# Pawn scale keys contributing to each planner STAT_COLUMNS weight (sum if multiple).
STAT_COL_PAWN_WEIGHT_KEYS = [
    ("Armor",),
    ("Strength",),
    ("Agility",),
    ("Stamina",),
    ("Intellect",),
    ("Spirit",),
    ("Healing",),
    ("SpellDamage",),
    ("Ap",),
    ("Mp5",),
    ("DefenseRating",),
    ("DodgeRating",),
    ("ParryRating",),
    ("BlockRating",),
    ("BlockValue",),
    ("HitRating",),
    ("CritRating",),
    ("SpellHitRating",),
    ("SpellCritRating",),
    ("HasteRating", "SpellHasteRating"),
    ("ResilienceRating",),
]

PAWN_GEMS_BC_LUA = os.path.join(SCRIPT_DIR, "Pawn", "GemsBurningCrusade.lua")

CSV_FIELDS = [
    "Spec", "Gear Type", "Phase", "Name", "Acquisition Type",
    "Quest", "Dungeon", "Difficulty", "Stats", "Notes",
]

SHEET_FIELDS = [
    "Interest", "Spec", "Gear Type", "Name", "Phase", "Acquisition Type",
    "Quest", "Dungeon", "Difficulty", "Stats",
    "Stat Comparison", "% Upgrade", "Notes",
]

# Ring/Trinket rows: no plain "Equipped" (use Equipped 1 / 2). Other rows: no Equipped 1/2.
INTEREST_OPTION_CLEAR = "----"
INTEREST_OPTIONS_RING_TRINKET = [
    INTEREST_OPTION_CLEAR,
    "Pass",
    "Consider",
    "Need",
    "Equipped 1",
    "Equipped 2",
]
INTEREST_OPTIONS_STANDARD = [
    INTEREST_OPTION_CLEAR,
    "Pass",
    "Consider",
    "Need",
    "Equipped",
]
INTEREST_OPTIONS = list(
    dict.fromkeys(INTEREST_OPTIONS_STANDARD + INTEREST_OPTIONS_RING_TRINKET)
)

SLOT_MAP = {
    "Head": "Head", "Shoulder": "Shoulder", "Back": "Back",
    "Chest": "Chest", "Wrist": "Wrist", "Hands": "Hands",
    "Waist": "Waist", "Legs": "Legs", "Feet": "Feet",
    "Neck": "Neck", "Ring": "Ring", "Trinket": "Trinket",
    "Main Hand": "Main Hand", "Off Hand": "Off Hand",
    "Two Hand": "Two Hand", "Ranged/Relic": "Ranged/Relic",
    "Main Hand~Off Hand": "Main Hand",
}

# Wowhead tooltip slot strings -> our slot names
TOOLTIP_SLOT_MAP = {
    "Head": "Head", "Shoulder": "Shoulder", "Back": "Back",
    "Chest": "Chest", "Robe": "Chest", "Wrist": "Wrist",
    "Hands": "Hands", "Waist": "Waist", "Legs": "Legs",
    "Feet": "Feet", "Neck": "Neck", "Finger": "Ring",
    "Trinket": "Trinket", "Main Hand": "Main Hand",
    "Off Hand": "Off Hand", "One-Hand": "Main Hand",
    "Two-Hand": "Two Hand", "Held In Off-hand": "Off Hand",
    "Ranged": "Ranged/Relic", "Relic": "Ranged/Relic",
    "Thrown": "Ranged/Relic", "Shield": "Off Hand",
}

SKIP_SOURCE_TYPES = {"Alchemy", "Transmute", "Conjured"}


def log(msg):
    print(msg, file=sys.stderr)


# ============================================================
# 1. PARSE ITEM SOURCES DB
# ============================================================

def parse_item_sources():
    filepath = os.path.join(DATA_DIR, "ItemSources.lua")
    sources = {}
    with open(filepath, "r") as f:
        for line in f:
            m = re.match(
                r'\s*\[(\d+)\]\s*=\s*\{\s*'
                r'Name\s*=\s*"([^"]*)".*?'
                r'SourceType\s*=\s*LBIS\.L\["([^"]*)"\].*?'
                r'Source\s*=\s*(.*?),\s*'
                r'SourceNumber\s*=\s*"([^"]*)".*?'
                r'SourceLocation\s*=\s*(.*?),\s*'
                r'SourceFaction\s*=\s*"([^"]*)"',
                line,
            )
            if not m:
                continue

            def extract_val(raw):
                lm = re.match(r'LBIS\.L\["([^"]*)"\]', raw)
                if lm:
                    return lm.group(1)
                sm = re.match(r'"([^"]*)"', raw)
                if sm:
                    return sm.group(1)
                if ".." in raw:
                    parts = re.findall(r'LBIS\.L\["([^"]*)"\]|"([^"]*)"', raw)
                    return " ".join(p[0] or p[1] for p in parts).strip()
                return raw.strip('"')

            sources[int(m.group(1))] = {
                "name": m.group(2),
                "source_type": m.group(3),
                "source": extract_val(m.group(4).strip()),
                "source_number": m.group(5),
                "source_location": extract_val(m.group(6).strip()),
                "faction": m.group(7),
            }
    return sources


# ============================================================
# 2. PARSE GUIDE LUA FILES
# ============================================================

def parse_guide_items(filepath, spec_name, max_phase):
    """Parse items from phases 0..max_phase (cumulative), plus phase 99 (PrePatch).

    Applies the same label transforms as the in-game addon:
      - Phase 0  -> "PreRaid"
      - Phase 99 -> "PrePatch"
      - Phase < max_phase -> replace "BIS" with "Alt"
      - Phase == max_phase -> keep original label
    Duplicates across phases use the highest phase's label.
    """
    phases_to_scan = list(range(max_phase + 1)) + [99]
    raw_by_phase = {p: [] for p in phases_to_scan}

    with open(filepath, "r") as f:
        content = f.read()

    for p in phases_to_scan:
        phase_var = f"spec{p}"
        for m in re.finditer(
            rf'LBIS:AddItem\({phase_var},\s*"(\d+)",\s*LBIS\.L\["([^"]*)"\],\s*"([^"]*)"\)\s*--(.+)',
            content,
        ):
            raw_by_phase[p].append({
                "item_id": int(m.group(1)),
                "gear_type": m.group(2),
                "bis_rank": m.group(3),
                "name": m.group(4).strip(),
            })

    seen = {}
    for p in phases_to_scan:
        for item in raw_by_phase[p]:
            if p == 0:
                label = "PreRaid"
            elif p == 99:
                label = "PrePatch"
            elif p < max_phase:
                label = item["bis_rank"].replace("BIS", "Alt")
            else:
                label = item["bis_rank"]

            key = (item["item_id"], item["gear_type"])
            if key in seen:
                if p != 99 and p > seen[key]["_phase"]:
                    seen[key]["bis_rank"] = label
                    seen[key]["_phase"] = p
            else:
                seen[key] = {
                    "spec": spec_name,
                    "item_id": item["item_id"],
                    "gear_type": item["gear_type"],
                    "bis_rank": label,
                    "name": item["name"],
                    "_phase": p,
                }

    return [
        {k: v for k, v in entry.items() if k != "_phase"}
        for entry in seen.values()
    ]


def parse_all_guide_items():
    """Parse all guide Lua files to get item_id -> slot mapping."""
    id_to_slot = {}
    id_to_name = {}
    for fname in os.listdir(LOON_GUIDES_DIR):
        if not fname.endswith(".lua"):
            continue
        path = os.path.join(LOON_GUIDES_DIR, fname)
        with open(path, "r") as f:
            for line in f:
                m = re.match(
                    r'LBIS:AddItem\(spec\d+,\s*"(\d+)",\s*LBIS\.L\["([^"]*)"\].*?\)\s*--(.+)',
                    line.strip(),
                )
                if m:
                    iid = int(m.group(1))
                    slot = SLOT_MAP.get(m.group(2), m.group(2))
                    id_to_slot.setdefault(iid, slot)
                    id_to_name[iid] = m.group(3).strip()
    return id_to_slot, id_to_name


# ============================================================
# 2b. CLASS / SPEC DISCOVERY
# ============================================================

def _title_case_words_apps_script_style(s: str) -> str:
    """Match apps_script.gs titleCaseWords — used when parsing CE lines like --- BEAR ---."""
    parts = str(s).split()
    out = []
    for w in parts:
        if not w:
            continue
        out.append(w[0].upper() + w[1:].lower())
    return " ".join(out)


def _validate_spec_order_for_ce_and_planner(spec_order: list) -> None:
    """
    CE section titles are f'--- {spec_name.upper()} ---'. Apps Script maps that back to column B
    via specDisplayNameFromSectionHeader → titleCaseWords(inner). That must equal the Loon/CSV spec
    string exactly, or Equip / N / O / Apps Script section matching breaks.
    """
    if len(spec_order) != len(set(spec_order)):
        log("ERROR: duplicate spec names in discovered spec_order")
        sys.exit(1)
    for spec in spec_order:
        canon = _title_case_words_apps_script_style(spec.upper())
        if canon != spec:
            log(
                "ERROR: Loon RegisterSpec second label %r must equal Apps Script "
                "titleCaseWords(upper(label)) → %r. Fix the lua string or rename to match."
                % (spec, canon)
            )
            sys.exit(1)


def _extract_spec_name(filepath):
    """Extract spec display name from the first RegisterSpec call in a guide file."""
    with open(filepath, "r") as f:
        for line in f:
            m = re.match(
                r'.*LBIS:RegisterSpec\(LBIS\.L\["[^"]*"\],\s*LBIS\.L\["([^"]*)"\]', line
            )
            if m:
                return m.group(1)
    return None


def discover_class_guides(class_name):
    """Auto-discover spec guides for a class by scanning AddonReference/Loon/.

    Returns (spec_order, guides_list) where guides_list is
    [(spec_name, filename), ...] sorted alphabetically by spec.
    """
    class_title = class_name.strip().title()

    guides = []
    for fname in sorted(os.listdir(LOON_GUIDES_DIR)):
        if fname.startswith(class_title) and fname.endswith(".lua"):
            path = os.path.join(LOON_GUIDES_DIR, fname)
            spec_name = _extract_spec_name(path)
            if spec_name:
                guides.append((spec_name, fname))

    if not guides:
        available = set()
        for fname in os.listdir(LOON_GUIDES_DIR):
            if fname.endswith(".lua"):
                m = re.match(r"([A-Z][a-z]+)", fname)
                if m:
                    available.add(m.group(1).lower())
        log(f"ERROR: No guide files found for class '{class_title}' in {LOON_GUIDES_DIR}")
        log(f"  Available classes: {', '.join(sorted(available))}")
        sys.exit(1)

    spec_order = [g[0] for g in guides]
    _validate_spec_order_for_ce_and_planner(spec_order)
    log(f"Discovered {class_title}: {', '.join(spec_order)} ({len(guides)} specs)")
    return spec_order, guides


def _detect_max_phase(filepath):
    """Detect the highest phase number registered in a guide file."""
    max_p = 0
    with open(filepath, "r") as f:
        for line in f:
            m = re.match(r'local\s+spec(\d+)\s*=\s*LBIS:RegisterSpec', line)
            if m:
                p = int(m.group(1))
                if p != 99 and p > max_p:
                    max_p = p
    return max_p


# ============================================================
# 3. TOOLTIP FETCHING + STAT PARSING
# ============================================================

def _make_ssl_ctx():
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return ctx


def fetch_tooltip(item_id, ssl_ctx, retries=3):
    url = f"https://nether.wowhead.com/tooltip/item/{item_id}?dataEnv={WOWHEAD_DATA_ENV}&locale=0"
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={
                "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)",
                "Accept": "application/json",
            })
            with urllib.request.urlopen(req, timeout=15, context=ssl_ctx) as resp:
                return json.loads(resp.read().decode("utf-8")).get("tooltip", "")
        except urllib.error.HTTPError as e:
            if e.code == 429:
                time.sleep(3 * (attempt + 1))
            elif e.code == 404:
                return None
            else:
                time.sleep(1)
        except Exception:
            time.sleep(1)
    return None


def _wowhead_tooltip_anchor_specials(html):
    """Equip/Use lines whose prose is only inside <a>…</a> (Wowhead JSON tooltips)."""
    out = []
    for label, pat in (
        ("Equip", r"Equip:\s*(?:<[^>]+>\s*)*<a[^>]*>([^<]{8,800})</a>"),
        ("Use", r"Use:\s*(?:<[^>]+>\s*)*<a[^>]*>([^<]{8,800})</a>"),
    ):
        for m in re.finditer(pat, html, re.I):
            text = m.group(1).strip().rstrip(".")
            if len(text) >= 8:
                out.append("%s: %s" % (label, text))
    return out


def _merge_anchor_specials(specials, html):
    """Append anchor-only effects without duplicating existing specials."""
    seen = {s.lower() for s in specials}
    for frag in _wowhead_tooltip_anchor_specials(html):
        fl = frag.lower()
        if fl not in seen:
            specials.append(frag)
            seen.add(fl)


def _parse_socket_counts(html):
    """Count gem sockets from Wowhead tooltip (class socket-meta / red / yellow / blue)."""
    counts = {"meta": 0, "red": 0, "yellow": 0, "blue": 0}
    if not html:
        return counts
    for m in re.finditer(r"\bsocket-(meta|red|yellow|blue)\b", html):
        counts[m.group(1)] += 1
    return counts


def _strip_set_bonus_blocks_from_tooltip_html(html):
    """Remove item set bonus lines from Wowhead tooltip HTML before stat extraction.

    Set bonuses use the same \"Increases your … rating\" phrasing as Equip: lines but only
    apply when multiple set pieces are worn; they must not be stored as base item stats or
    included in gear score (same idea as Pawn ignoring set bonuses).
    Typical markup: (2) Set : <a href=\"…\">Increases your hit rating by 35.</a>
    """
    if not html:
        return html
    return re.sub(
        r"\(\d+\)\s*Set\s*:\s*(?:<[^>]+>\s*)*<a[^>]*>.*?</a>",
        "",
        html,
        flags=re.I | re.DOTALL,
    )


def parse_tooltip_to_dict(html):
    """Parse tooltip HTML into (stats_dict, special_str, slot_str, item_level, sockets_dict)."""
    if not html:
        return {}, "", None, 0, _parse_socket_counts("")

    # Strip Wowhead <!--…--> placeholders (e.g. <!--rtg32-->) so "+rating by 32" patterns match.
    html = re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)
    html = _strip_set_bonus_blocks_from_tooltip_html(html)

    stats = {}

    armor_m = re.search(r"(\d+)\s+Armor", html)
    if armor_m:
        stats["Armor"] = int(armor_m.group(1))

    dps_m = re.search(r"([\d.]+) damage per second", html, re.I)
    if dps_m:
        stats["DPS"] = float(dps_m.group(1))

    for pattern, key in [
        (r"\+(\d+)\s+Strength", "Str"),
        (r"\+(\d+)\s+Agility", "Agi"),
        (r"\+(\d+)\s+Stamina", "Sta"),
        (r"\+(\d+)\s+Intellect", "Int"),
        (r"\+(\d+)\s+Spirit", "Spi"),
    ]:
        m = re.search(pattern, html)
        if m:
            stats[key] = int(m.group(1))

    equip_patterns = [
        (r"[Ii](?:ncreases?|mproves?) (?:your )?defense rating by\s*(\d+)", "Defense"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?dodge rating by\s*(\d+)", "Dodge"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?parry rating by\s*(\d+)", "Parry"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?(?:shield )?block rating by\s*(\d+)", "Block Rating"),
        (r"[Ii](?:ncreases?|mproves?) the block value of your shield by\s*(\d+)", "Block Value"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?resilience rating by\s*(\d+)", "Resilience"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?(?:melee and ranged )?hit rating by\s*(\d+)", "Hit"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?(?:melee and ranged )?critical strike rating by\s*(\d+)", "Crit"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?spell hit rating by\s*(\d+)", "Spell Hit"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?spell critical strike rating by\s*(\d+)", "Spell Crit"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?(?:melee )?haste rating by\s*(\d+)", "Haste"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?expertise rating by\s*(\d+)", "Expertise"),
        (r"[Ii](?:ncreases?|mproves?) (?:your )?armor penetration rating by\s*(\d+)", "ArPen"),
        (r"[Ii](?:ncreases?|mproves?) attack power by\s*(\d+)", "AP"),
        (r"[Ii](?:ncreases?|mproves?) damage and healing done by magical spells and effects by (?:up to )?(\d+)", "Spell Dmg/Heal"),
        (r"[Rr]estores? (\d+) mana per 5 sec", "MP5"),
        (r"[Rr]estores? (\d+) health per 5 sec", "HP5"),
    ]
    for pattern, key in equip_patterns:
        m = re.search(pattern, html)
        if m:
            stats[key] = int(m.group(1))

    heal_m = re.search(r"[Ii]ncreases? healing done by (?:up to )?(\d+)(?:.*?damage done by (?:up to )?(\d+))?", html)
    if heal_m:
        stats["Healing"] = int(heal_m.group(1))
        if heal_m.lastindex >= 2 and heal_m.group(2):
            stats["Spell Dmg"] = int(heal_m.group(2))

    specials = []
    use_m = re.search(r"Use:\s*(?:<[^>]*>)*\s*([^<]+?)(?:\.\s*\(|<)", html)
    if use_m:
        eff = use_m.group(1).strip().rstrip(".")
        if 5 < len(eff) < 200:
            specials.append(f"Use: {eff}")

    proc_m = re.search(r"Chance on (?:hit|melee hit|spell critical hit|critical hit):\s*(?:<[^>]*>)*\s*([^<]+?)(?:\.\s*|<)", html)
    if proc_m:
        eff = proc_m.group(1).strip().rstrip(".")
        if len(eff) > 5:
            specials.append(f"Proc: {eff}")

    equip_special = re.findall(r"Equip:\s*(?:<[^>]*>)*\s*([^<]+?)(?:\.|<)", html)
    for eff in equip_special:
        eff_clean = eff.strip().rstrip(".")
        already = any(
            kw in eff_clean.lower()
            for kw in [
                "defense", "dodge", "parry", "block", "resilience", "hit",
                "critical", "haste", "expertise", "armor penetration",
                "attack power", "damage and healing", "healing done",
                "mana per 5", "health per 5", "damage done",
            ]
        )
        if not already and 10 < len(eff_clean) < 200:
            specials.append(f"Equip: {eff_clean}")

    _merge_anchor_specials(specials, html)

    special_str = "; ".join(specials).replace("&nbsp;", " ").replace("  ", " ").strip()

    slot = None
    slot_m = re.search(r'<table[^>]*><tr><td>([A-Za-z][A-Za-z -]+?)</td>', html)
    if slot_m:
        raw_slot = slot_m.group(1).strip()
        slot = TOOLTIP_SLOT_MAP.get(raw_slot)

    ilvl = 0
    ilvl_m = re.search(r"Item Level\s*(?:<!--[^>]*-->)?\s*(\d+)", html)
    if ilvl_m:
        ilvl = int(ilvl_m.group(1))

    normalize_spell_stats(stats)
    special_str = strip_equip_redundant_with_display_stats(stats, special_str)

    sockets = _parse_socket_counts(html)
    return stats, special_str, slot, ilvl, sockets


def normalize_spell_stats(stats):
    """Map spell tooltip data to Healing + Spell Dmg only (no stored Spell Dmg/Heal key).

    Parse may still set Spell Dmg/Heal from combined green text; split lines from heal_m win on
    conflict. Combined-only tooltips become equal Healing and Spell Dmg; _spell_from_combined_green
    marks that case for idempotent reloads.
    """
    if not stats:
        return

    if "Spell Dmg/Heal" in stats and (
        "Healing" in stats or "Spell Dmg" in stats
    ):
        stats.pop("_spell_from_combined_green", None)
        del stats["Spell Dmg/Heal"]

    if "Spell Dmg/Heal" in stats and "Healing" not in stats and "Spell Dmg" not in stats:
        v = stats["Spell Dmg/Heal"]
        stats["Healing"] = v
        stats["Spell Dmg"] = v
        del stats["Spell Dmg/Heal"]
        stats["_spell_from_combined_green"] = True

    stats.pop("Spell Dmg/Heal", None)


def _display_stats_for_equip_filter(stats):
    """Subset of stats that appear in the planner Stats column (STAT_COLUMNS only)."""
    if not stats:
        return {}
    out = {}
    for k in STAT_COLUMNS:
        v = stats.get(k)
        if v is None or v == "" or isinstance(v, bool):
            continue
        try:
            out[k] = int(v)
        except (TypeError, ValueError):
            continue
    return out


def _split_special_segments(special_str):
    """Split Use:/Equip:/Proc: clauses joined with '; ' without breaking prose semicolons."""
    s = (special_str or "").strip()
    if not s:
        return []
    parts = re.split(r";\s+(?=Use:|Equip:|Proc:)", s, flags=re.I)
    return [p.strip() for p in parts if p.strip()]


def _equip_body_is_proc_like(body):
    bl = body.lower()
    return (
        "chance on" in bl
        or "proc chance" in bl
        or "have a chance" in bl
        or "chance to " in bl
        or " cooldown" in bl
        or "cooldown)" in bl
    )


def _equip_segment_redundant_with_display_stats(display_stats, seg):
    """True if this Equip: line only repeats numbers already shown in STAT_COLUMNS."""
    m = re.match(r"^Equip:\s*(.*)$", seg.strip(), re.I | re.DOTALL)
    if not m:
        return False
    body = m.group(1).strip()
    if not body or _equip_body_is_proc_like(body):
        return False

    ds = display_stats

    def n_same(key, n):
        v = ds.get(key)
        return v is not None and int(n) == int(v)

    # --- Order: more specific patterns before looser ones ---

    m2 = re.search(
        r"[Ii](?:ncreases?|mproves?)\s+damage\s+and\s+healing\s+done\s+by\s+magical\s+spells\s+and\s+effects\s+by\s+(?:up\s+to\s+)?(\d+)",
        body,
    )
    if m2:
        n = int(m2.group(1))
        return n_same("Healing", n) and n_same("Spell Dmg", n)

    m2 = re.search(
        r"[Ii]ncreases?\s+healing\s+done\s+by\s+(?:up\s+to\s+)?(\d+)(?:.*?damage\s+done\s+by\s+(?:up\s+to\s+)?(\d+))?",
        body,
    )
    if m2:
        h = int(m2.group(1))
        if m2.lastindex >= 2 and m2.group(2):
            d = int(m2.group(2))
            return n_same("Healing", h) and n_same("Spell Dmg", d)
        return n_same("Healing", h)

    m2 = re.search(
        r"[Ii](?:ncreases?|mproves?)\s+damage\s+done\s+by\s+magical\s+spells\s+and\s+effects\s+by\s+(?:up\s+to\s+)?(\d+)",
        body,
    )
    if m2:
        n = int(m2.group(1))
        return n_same("Spell Dmg", n)

    m2 = re.search(
        r"[Ii](?:ncreases?|mproves?)\s+attack\s+power\s+by\s+(\d+)", body
    )
    if m2:
        return n_same("AP", int(m2.group(1)))

    rating_specs = [
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?defense\s+rating\s+by\s*(\d+)",
            "Defense",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?dodge\s+rating\s+by\s*(\d+)",
            "Dodge",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?parry\s+rating\s+by\s*(\d+)",
            "Parry",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?(?:shield\s+)?block\s+rating\s+by\s*(\d+)",
            "Block Rating",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+the\s+block\s+value\s+of\s+your\s+shield\s+by\s*(\d+)",
            "Block Value",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?(?:melee\s+and\s+ranged\s+)?hit\s+rating\s+by\s*(\d+)",
            "Hit",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?(?:melee\s+and\s+ranged\s+)?critical\s+strike\s+rating\s+by\s*(\d+)",
            "Crit",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?spell\s+hit\s+rating\s+by\s*(\d+)",
            "Spell Hit",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?spell\s+critical\s+strike\s+rating\s+by\s*(\d+)",
            "Spell Crit",
        ),
        (
            r"[Ii](?:ncreases?|mproves?)\s+(?:your\s+)?(?:melee\s+)?haste\s+rating\s+by\s*(\d+)",
            "Haste",
        ),
    ]
    for pat, key in rating_specs:
        m2 = re.search(pat, body)
        if m2 and key in STAT_COLUMNS:
            return n_same(key, int(m2.group(1)))

    m2 = re.search(r"[Rr]estores?\s+(\d+)\s+mana\s+per\s+5\s+sec", body)
    if m2:
        return n_same("MP5", int(m2.group(1)))

    return False


def strip_equip_redundant_with_display_stats(stats, special_str):
    """Drop Equip: segments whose numeric effects are already shown in STAT_COLUMNS."""
    if not special_str or not str(special_str).strip():
        return (special_str or "").strip()
    ds = _display_stats_for_equip_filter(stats)
    kept = []
    for seg in _split_special_segments(special_str):
        if re.match(r"^Equip:", seg.strip(), re.I):
            if _equip_segment_redundant_with_display_stats(ds, seg):
                continue
        kept.append(seg)
    return "; ".join(kept).strip()


def _break_use_equip_newlines(text: str) -> str:
    """Ensure space-prefixed Use:/Equip: in prose start on new lines (Stats / display)."""
    if not text or not str(text).strip():
        return (text or "").strip()
    t = str(text).strip()
    t = t.replace(" Use:", "\nUse:")
    t = t.replace(" Equip:", "\nEquip:")
    return t


def stats_dict_to_string(stats, attributes=""):
    """Convert a stats dict + attribute string to the readable Stats column.

    Stat lines use '; ' between pairs; two newlines separate the stat block from
    on-use / equip attribute text when both are present.
    """
    parts = []
    for key in STAT_COLUMNS:
        val = stats.get(key)
        if val:
            parts.append(f"{val} {key}")
    result = "; ".join(parts)
    if attributes:
        at = strip_equip_redundant_with_display_stats(stats, attributes.strip())
        at = _break_use_equip_newlines(at)
        if at:
            if result:
                result += "\n\n" + at
            else:
                result = at
    return result


# ============================================================
# 4. BUILD ITEM DATABASE (--build-db)
# ============================================================

_fetch_counter_lock = threading.Lock()
_fetch_counter = [0, 0]  # [done, total]


def _fetch_one_item(item_id, name, ssl_ctx):
    tooltip = fetch_tooltip(item_id, ssl_ctx)
    with _fetch_counter_lock:
        _fetch_counter[0] += 1
        done, total = _fetch_counter
        if done % 50 == 0 or done == total:
            log(f"  [{done}/{total}] ...")
    time.sleep(FETCH_DELAY)
    if not tooltip:
        return item_id, None
    stats, special, slot, ilvl, sockets = parse_tooltip_to_dict(tooltip)
    return item_id, {
        "name": name,
        "slot": slot,
        "stats": stats,
        "special": special,
        "ilvl": ilvl,
        "sockets": sockets,
    }


def _cache_entry_is_stale_tooltip_(entry):
    """True if cached item likely predates current tooltip parser (empty stats + empty special but ilvl set)."""
    if not entry or not isinstance(entry, dict):
        return False
    stats = entry.get("stats") or {}
    special = (entry.get("special") or "").strip()
    if stats or special:
        return False
    return int(entry.get("ilvl") or 0) > 0


def build_item_database(test_mode=False, test_class=None, force_refresh=False):
    existing_db = {}
    if os.path.exists(ITEM_DB_PATH):
        with open(ITEM_DB_PATH, "r") as f:
            existing_db = json.load(f)
        log(f"Loaded existing database: {len(existing_db)} items")

    guide_slots, guide_names = parse_all_guide_items()
    log(f"Parsed guide files: {len(guide_slots)} unique items with known slots")

    sources = parse_item_sources()
    log(f"Parsed ItemSources.lua: {len(sources)} items")

    if test_mode:
        test_ids = set()
        if test_class:
            _, guides = discover_class_guides(test_class)
        else:
            guides = [
                (None, f)
                for f in os.listdir(LOON_GUIDES_DIR) if f.endswith(".lua")
            ]
        for _, filename in guides:
            path = os.path.join(LOON_GUIDES_DIR, filename)
            if os.path.exists(path):
                with open(path, "r") as f:
                    for line in f:
                        m = re.match(r'LBIS:AddItem\(spec\d+,\s*"(\d+)"', line.strip())
                        if m:
                            test_ids.add(int(m.group(1)))
        label = test_class.title() if test_class else "all guide"
        items_to_fetch = {
            iid: guide_names.get(iid, sources.get(iid, {}).get("name", f"Item {iid}"))
            for iid in test_ids
        }
        log(f"TEST MODE: only fetching {len(items_to_fetch)} {label} items")
    else:
        items_to_fetch = {}
        for iid, info in sources.items():
            if info["source_type"] in SKIP_SOURCE_TYPES:
                continue
            src = info.get("source", "")
            if info["source_type"] == "Profession" and any(
                skip in src for skip in ("Alchemy", "Cooking", "First Aid")
            ):
                continue
            items_to_fetch[iid] = info["name"]
        for iid, name in guide_names.items():
            if iid not in items_to_fetch:
                items_to_fetch[iid] = name

    to_fetch = {}
    n_new = 0
    n_stale = 0
    for iid, name in items_to_fetch.items():
        sid = str(iid)
        if sid not in existing_db:
            to_fetch[iid] = name
            n_new += 1
        elif force_refresh:
            to_fetch[iid] = name
        elif _cache_entry_is_stale_tooltip_(existing_db.get(sid)):
            to_fetch[iid] = name
            n_stale += 1
    n_skip = len(items_to_fetch) - len(to_fetch)
    if force_refresh:
        log(f"Items to fetch from Wowhead: {len(to_fetch)} (--refresh-db, full catalog refetch)")
    else:
        log(
            f"Items to fetch from Wowhead: {len(to_fetch)} "
            f"({n_new} new, {n_stale} stale empty-tooltip, {n_skip} cached ok)"
        )

    if to_fetch:
        _fetch_counter[0] = 0
        _fetch_counter[1] = len(to_fetch)
        ssl_ctx = _make_ssl_ctx()

        with concurrent.futures.ThreadPoolExecutor(max_workers=FETCH_WORKERS) as pool:
            futures = {
                pool.submit(_fetch_one_item, iid, name, ssl_ctx): iid
                for iid, name in to_fetch.items()
            }
            for future in concurrent.futures.as_completed(futures):
                iid, result = future.result()
                if result:
                    existing_db[str(iid)] = result

    db = {}
    for str_id, entry in existing_db.items():
        iid = int(str_id)
        slot = entry.get("slot")
        if not slot and iid in guide_slots:
            entry["slot"] = guide_slots[iid]
            slot = entry["slot"]
        if slot:
            db[str_id] = entry

    for entry in db.values():
        normalize_spell_stats(entry.get("stats") or {})

    os.makedirs(DATA_DIR, exist_ok=True)
    with open(ITEM_DB_PATH, "w") as f:
        json.dump(db, f, separators=(",", ":"))

    log(f"\nSaved {len(db)} gear items -> {ITEM_DB_PATH}")
    return db


def load_item_database():
    if not os.path.exists(ITEM_DB_PATH):
        log(f"ERROR: {ITEM_DB_PATH} not found. Run --build-db first.")
        sys.exit(1)
    with open(ITEM_DB_PATH, "r") as f:
        data = json.load(f)
    for entry in data.values():
        if isinstance(entry.get("stats"), dict):
            normalize_spell_stats(entry["stats"])
    return data


# ============================================================
# 5. CLASSIFY ACQUISITION
# ============================================================

def _abbrev_location(loc):
    if not loc:
        return loc
    clean = loc.replace(" (H)", "").replace(" (N)", "").replace(" (Heroic)", "").replace(" (Normal)", "")
    return LOCATION_ABBREVS.get(clean, clean)


def _is_dungeon_quest(loc):
    if not loc:
        return False
    return any(kw.lower() in loc.lower() for kw in DUNGEON_KEYWORDS)


def _shorten_dungeon(name):
    return name[4:] if name.startswith("The ") else name


def classify_acquisition(src_info):
    """Returns (acq_type, quest, dungeon, difficulty, notes)."""
    if not src_info:
        return "Other", "", "", "", ""

    st = src_info["source_type"]
    source = src_info["source"]
    loc = src_info["source_location"]

    if st == "Profession" or st == "Transmute":
        return "Auction House", "", "", "", ""

    if st == "Quest":
        abbr = _abbrev_location(loc)
        kind = "Quest (Dung)" if _is_dungeon_quest(loc) else "Quest"
        return kind, f"{source} ({abbr})", "", "", ""

    if st == "Drop":
        if source == "World Drop" or not loc:
            return "Auction House", "", "", "", "World Drop"
        if loc in ("World Bosses", "World Boss"):
            return "World Drop", "", "", "", source
        is_heroic = any(h in loc for h in ("(H)", "(Heroic)"))
        clean_loc = re.sub(r"\s*\((H|N|Heroic|Normal)\)", "", loc)
        dungeon = _shorten_dungeon(clean_loc)
        difficulty = "Heroic" if is_heroic else "Normal"
        return "Dungeon Drop", "", dungeon, difficulty, source

    if st == "Dungeon Token":
        cost = src_info["source_number"]
        if "Badge of Justice" in source:
            note = f"{cost} Badges of Justice"
        elif "Spirit Shard" in source:
            note = f"{cost} Spirit Shards"
        else:
            note = f"{source} x{cost}"
        return "Dungeon Token", "", "", "", note

    if st == "Reputation":
        return "Reputation", "", "", "", f"{source} - {loc}"

    if st == "PvP":
        return "PvP", "", "", "", ""

    if st == "Vendor":
        return "Vendor", "", "", "", loc

    if st == "Tier Token":
        return "Raid Drop", "", "", "", f"{source} ({loc})"

    if st == "Token":
        return "Dungeon Token", "", "", "", source

    if st == "Conjured":
        return "Other", "", "", "", source

    return "Other", "", "", "", ""


# ============================================================
# 6. GENERATE CSV
# ============================================================

def generate_csv(max_phase, class_name, output_csv):
    spec_order, guides = discover_class_guides(class_name)

    item_db = load_item_database()
    id_to_stats = {}
    for str_id, entry in item_db.items():
        id_to_stats[int(str_id)] = entry

    sources = parse_item_sources()
    log(f"Parsed ItemSources.lua: {len(sources)} items")

    all_items = []
    for spec_name, filename in guides:
        path = os.path.join(LOON_GUIDES_DIR, filename)
        if not os.path.exists(path):
            log(f"  WARNING: {path} not found, skipping")
            continue
        items = parse_guide_items(path, spec_name, max_phase)
        all_items.extend(items)
        log(f"  {spec_name}: {len(items)} items (phases 0-{max_phase})")

    rows = []
    for item in all_items:
        src = sources.get(item["item_id"])
        acq_type, quest, dungeon, difficulty, notes = classify_acquisition(src)
        db_entry = id_to_stats.get(item["item_id"], {})
        stats_d = db_entry.get("stats", {})
        attributes = db_entry.get("special", "")
        stats_str = stats_dict_to_string(stats_d, attributes)
        rows.append({
            "Spec": item["spec"],
            "Gear Type": item["gear_type"],
            "Phase": item["bis_rank"],
            "Name": item["name"],
            "Acquisition Type": acq_type,
            "Quest": quest,
            "Dungeon": dungeon,
            "Difficulty": difficulty,
            "Stats": stats_str,
            "Notes": notes,
        })

    sort_key = _make_sort_key(spec_order)
    rows.sort(key=sort_key)

    os.makedirs(os.path.dirname(output_csv), exist_ok=True)
    with open(output_csv, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    log(f"Wrote {len(rows)} rows -> {output_csv}")
    return rows, spec_order


def _make_sort_key(spec_order):
    def _sort_key(row):
        s = spec_order.index(row["Spec"]) if row["Spec"] in spec_order else 99
        g = GEAR_ORDER.index(row["Gear Type"]) if row["Gear Type"] in GEAR_ORDER else 99
        a = ACQ_ORDER.index(row["Acquisition Type"]) if row["Acquisition Type"] in ACQ_ORDER else 99
        return (s, g, a, row["Name"])
    return _sort_key


# ============================================================
# 7. EXCEL EXPORT
# ============================================================

ACQ_COLORS = {
    "Auction House": "80DEEA", "Reputation": "A5D6A7",
    "Quest": "FFE082", "Quest (Dung)": "FFE082",
    "PvP": "F48FB1", "Dungeon Token": "BDBDBD",
    "World Drop": "EEEEEE", "Raid Drop": "D7CCC8",
    "Vendor": "EEEEEE", "Other": "EEEEEE",
}

DUNGEON_COLORS_NORMAL = {
    "Blood Furnace": "FFD5D5", "Hellfire Ramparts": "FFE0CC",
    "Shattered Halls": "FFEABB", "Mechanar": "FFF5BB",
    "Botanica": "E8FFCC", "Arcatraz": "FFF0CC",
    "Auchenai Crypts": "F5CCF0", "Mana Tombs": "D5CCF5",
    "Sethekk Halls": "E0CCF5", "Shadow Labyrinth": "E8CCEE",
    "Slave Pens": "CCE5FF", "Underbog": "CCEED8",
    "Steamvault": "CCD5F5",
    "Old Hillsbrad Foothills": "EEE4CC", "Black Morass": "CCE8D8",
    "Dire Maul": "F0EADD", "Molten Core": "FFE0D0",
    "Blackwing Lair": "D8DEE8", "Ahn'Qiraj": "F5F0CC",
    "Naxxramas": "D0EEEE", "Terokkar Forest": "DEE8CC",
}
DUNGEON_COLORS_HEROIC = {
    "Blood Furnace": "E87070", "Hellfire Ramparts": "E89060",
    "Shattered Halls": "D8A040", "Mechanar": "C8B030",
    "Botanica": "70B848", "Arcatraz": "D89838",
    "Auchenai Crypts": "C850A8", "Mana Tombs": "7858C8",
    "Sethekk Halls": "9858C8", "Shadow Labyrinth": "A848B8",
    "Slave Pens": "4090D8", "Underbog": "40A878",
    "Steamvault": "4868C8",
    "Old Hillsbrad Foothills": "B89050", "Black Morass": "48A880",
    "Dire Maul": "A89868", "Molten Core": "D07050",
    "Blackwing Lair": "7888A8", "Ahn'Qiraj": "B0A040",
    "Naxxramas": "48A8B0", "Terokkar Forest": "80A848",
}

# WoW class IDs (retail API) — matches Pawn ClassicHawsJon.lua (no DK in TBC).
CLASS_TITLE_TO_PAWN_CLASS_ID = {
    "Warrior": 1,
    "Paladin": 2,
    "Hunter": 3,
    "Rogue": 4,
    "Priest": 5,
    "Shaman": 7,
    "Mage": 8,
    "Warlock": 9,
    "Druid": 11,
}

# Loon `RegisterSpec(..., LBIS.L["SpecLabel"], ...)` vs Pawn `spec_name` in data/pawn_scales_tbc.json.
# Known mismatches (anything else must match Pawn exactly):
#   Druid Cat/Bear -> Feral (Damage) / Feral (Tank)
#   Rogue: one Loon guide "Dps" vs Pawn Assassination / Combat / Subtlety -> map to Combat (TBC default)
#   Priest: Pawn has Discipline; Loon only Holy + Shadow (no Discipline guide / CE section)
PLANNER_SPEC_TO_PAWN_SPEC_NAME = {
    ("Druid", "Cat"): "Feral (Damage)",
    ("Druid", "Bear"): "Feral (Tank)",
    ("Rogue", "Dps"): "Combat",
}

_PAWN_VARS_TBC = {
    "HitRatingPer": 1.0,
    "SpellHitRatingPer": 1.0,
    "CritRatingPer": 1.0,
    "SpellCritRatingPer": 1.0,
    "HasteRatingPer": 1.0,
    "SpellHasteRatingPer": 1.0,
    "ExpertiseRatingPer": 1.0,
    "ArmorPenetrationPer": 1.0,
    "SpellPenetrationPer": 1.0,
    "DefenseRatingPer": 1.0,
    "DodgeRatingPer": 1.0,
    "ParryRatingPer": 1.0,
    "BlockRatingPer": 1.0,
}


def _pawn_eval_rhs(rhs):
    rhs = (rhs or "").strip()
    if rhs == "PawnIgnoreStatValue":
        return None
    if "*" in rhs:
        left, right = rhs.split("*", 1)
        left, right = left.strip(), right.strip()
        try:
            mult = float(_PAWN_VARS_TBC[left])
        except KeyError:
            mult = 1.0
        return mult * float(right)
    try:
        return float(rhs)
    except ValueError:
        return None


def _pawn_extract_table_body(s):
    start = s.find("{")
    if start < 0:
        return None
    depth = 0
    for i in range(start, len(s)):
        if s[i] == "{":
            depth += 1
        elif s[i] == "}":
            depth -= 1
            if depth == 0:
                return s[start + 1 : i]
    return None


def _pawn_parse_weights(body):
    out = {}
    for p in re.split(r",\s*", body):
        p = p.strip()
        if not p:
            continue
        m = re.match(r"^(\w+)\s*=\s*(.+)$", p)
        if not m:
            continue
        val = _pawn_eval_rhs(m.group(2))
        if val is not None:
            out[m.group(1)] = val
    return out


def _pawn_extract_scales_from_lua(lua_text):
    """Parse Classic + TBC branch from ClassicHawsJon.lua into scale dicts."""
    try:
        i0 = lua_text.index("if VgerCore.IsClassic or VgerCore.IsBurningCrusade then")
        i1 = lua_text.index("elseif VgerCore.IsWrath", i0)
    except ValueError:
        return None
    section = lua_text[i0:i1]
    chunks = re.split(r"PawnAddPluginScaleFromTemplate\s*\(\s*", section)
    scales = []
    for raw in chunks[1:]:
        lines = raw.split("\n")
        idx = 0
        while idx < len(lines) and "ScaleProviderName" not in lines[idx]:
            idx += 1
        if idx >= len(lines):
            continue
        idx += 1
        while idx < len(lines) and not re.search(r"^\s*\d+\s*,\s*--", lines[idx]):
            idx += 1
        if idx >= len(lines):
            continue
        m_cls = re.search(r"^\s*(\d+)\s*,\s*--\s*(.*)$", lines[idx])
        if not m_cls:
            continue
        class_id = int(m_cls.group(1))
        class_name = m_cls.group(2).strip()
        idx += 1
        m_spec = re.search(r"^\s*(?:(\d+)|nil)\s*,\s*(?:--\s*(.*))?$", lines[idx])
        if not m_spec or m_spec.group(1) is None:
            continue
        spec_id = int(m_spec.group(1))
        spec_name = (m_spec.group(2) or "").strip()
        rest = "\n".join(lines[idx:])
        body = _pawn_extract_table_body(rest)
        if not body:
            continue
        weights = _pawn_parse_weights(body)
        scales.append(
            {
                "class_id": class_id,
                "class_name": class_name,
                "spec_id": spec_id,
                "spec_name": spec_name,
                "weights": weights,
            }
        )
    return scales


def planner_weight_from_pawn_keys(weights, keys):
    """Sum Pawn scale weights for one planner stat column."""
    if not keys:
        return 0.0
    t = 0.0
    for k in keys:
        v = weights.get(k)
        if v is not None:
            t += float(v)
    return t


def pawn_weights_vector_for_stat_columns(weights_dict):
    """List of len(STAT_COLUMNS) floats for Current Equipment weight row."""
    return [
        planner_weight_from_pawn_keys(weights_dict, keys)
        for keys in STAT_COL_PAWN_WEIGHT_KEYS
    ]


def get_pawn_weights_for_spec(class_title, spec_name):
    """Return (weights_dict or None, vector aligned to STAT_COLUMNS)."""
    cid = CLASS_TITLE_TO_PAWN_CLASS_ID.get((class_title or "").strip().title())
    if cid is None or not os.path.isfile(PAWN_SCALES_JSON_PATH):
        return None, None
    try:
        with open(PAWN_SCALES_JSON_PATH, "r", encoding="utf-8") as f:
            payload = json.load(f)
    except (OSError, json.JSONDecodeError):
        return None, None
    cls_title = (class_title or "").strip().title()
    want = (spec_name or "").strip()
    want_pawn = PLANNER_SPEC_TO_PAWN_SPEC_NAME.get((cls_title, want), want)
    for sc in payload.get("scales", []):
        if sc.get("class_id") != cid:
            continue
        if (sc.get("spec_name") or "").strip() != want_pawn:
            continue
        wd = sc.get("weights") or {}
        return wd, pawn_weights_vector_for_stat_columns(wd)
    return None, None


def _lua_skip_ws(s, i):
    while i < len(s) and s[i] in " \t\n\r":
        i += 1
    return i


def _lua_brace_span(s, open_idx):
    """Return (inner_without_braces, index_after_closing) or (None, open_idx)."""
    if open_idx >= len(s) or s[open_idx] != "{":
        return None, open_idx
    depth = 0
    for j in range(open_idx, len(s)):
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[open_idx + 1 : j], j + 1
    return None, open_idx


def _lua_parse_stats_table(inner):
    out = {}
    if not inner:
        return out
    parts = re.split(r",\s*", inner.strip())
    for p in parts:
        p = p.strip()
        if not p:
            continue
        m = re.match(r"^(\w+)\s*=\s*([\d.]+)\s*$", p)
        if m:
            v = m.group(2)
            out[m.group(1)] = float(v) if "." in v else int(float(v))
    return out


def _gem_categorize(has_r, has_y, has_b, is_meta_table):
    if is_meta_table:
        return "meta"
    if has_r and has_y and not has_b:
        return "orange"
    if has_r and has_b and not has_y:
        return "purple"
    if has_y and has_b and not has_r:
        return "green"
    if has_r and not has_y and not has_b:
        return "red"
    if has_y and not has_r and not has_b:
        return "yellow"
    if has_b and not has_r and not has_y:
        return "blue"
    return None


def _pawn_stat_display_label(pawn_key, val):
    """Short label for gem dropdown (+8 Str style)."""
    short = {
        "Strength": "Str",
        "Agility": "Agi",
        "Stamina": "Sta",
        "Intellect": "Int",
        "Spirit": "Spi",
        "SpellDamage": "Spell Dmg",
        "Healing": "Healing",
        "Ap": "AP",
        "Rap": "RAP",
        "Armor": "Armor",
        "Mp5": "MP5",
        "HitRating": "Hit",
        "CritRating": "Crit",
        "SpellHitRating": "Spell Hit",
        "SpellCritRating": "Spell Crit",
        "HasteRating": "Haste",
        "SpellHasteRating": "Spell Haste",
        "DefenseRating": "Def",
        "DodgeRating": "Dodge",
        "ParryRating": "Parry",
        "BlockRating": "Block",
        "BlockValue": "BV",
        "SpellPenetration": "Spell Pen",
        "ExpertiseRating": "Exp",
        "ResilienceRating": "Res",
    }.get(pawn_key, pawn_key)
    iv = int(val) if float(val) == int(float(val)) else val
    return "+%s %s" % (iv, short)


def format_gem_stats_label(stats_dict):
    """Sorted deterministic label for dedupe / dropdown."""
    if not stats_dict:
        return ""
    items = sorted(stats_dict.items(), key=lambda z: z[0].lower())
    return "/".join(_pawn_stat_display_label(k, v) for k, v in items)


# Google Sheets treats leading "+" like a formula; prefix ZWSP so the cell is plain text.
GEM_LABEL_SHEETS_PREFIX = "\u200b"


def _gem_sort_key_label(label: str) -> tuple:
    """Sort: group by alphabetically-first stat token, then descending max bonus, then label."""
    s = (label or "").replace(GEM_LABEL_SHEETS_PREFIX, "").strip()
    parts = re.findall(r"\+?(\d+(?:\.\d+)?)\s+([^/]+)", s)
    if not parts:
        return ("zzz", 0, s.lower())
    stat_names = [p[1].strip().lower() for p in parts]
    bucket = min(stat_names)
    max_bonus = max(int(float(p[0])) for p in parts)
    return (bucket, -max_bonus, s.lower())


def _finalize_gem_label_for_sheet(label: str) -> str:
    if not label:
        return label
    if label.startswith(GEM_LABEL_SHEETS_PREFIX):
        return label
    return GEM_LABEL_SHEETS_PREFIX + label


def parse_pawn_gems_burning_crusade(lua_path=None):
    """
    Parse Pawn/GemsBurningCrusade.lua into categories:
    red, yellow, blue, orange, purple, green, meta.
    Each entry: { "id", "category", "stats", "label" }.
    """
    path = lua_path or PAWN_GEMS_BC_LUA
    if not os.path.isfile(path):
        return {c: [] for c in ("red", "yellow", "blue", "orange", "purple", "green", "meta")}
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()

    table_names = [
        ("PawnGemData60Common", False),
        ("PawnGemData70Uncommon", False),
        ("PawnGemData70Rare", False),
        ("PawnGemData70Epic", False),
        ("PawnMetaGemData70Rare", True),
    ]
    by_cat = {c: [] for c in ("red", "yellow", "blue", "orange", "purple", "green", "meta")}
    seen = set()

    for tbl, is_meta in table_names:
        m = re.search(
            r"local\s+" + re.escape(tbl) + r"\s*=\s*\{", text
        )
        if not m:
            continue
        open_idx = m.end() - 1
        inner, _ = _lua_brace_span(text, open_idx)
        if inner is None:
            continue
        pos = 0
        while pos < len(inner):
            j = inner.find("{ ID =", pos)
            if j < 0:
                break
            blk_inner, after = _lua_brace_span(inner, j)
            if blk_inner is None:
                pos = j + 1
                continue
            pos = after
            id_m = re.search(r"ID\s*=\s*(\d+)", blk_inner)
            if not id_m:
                continue
            gid = int(id_m.group(1))
            has_r = bool(re.search(r"\bR\s*=\s*true\b", blk_inner))
            has_y = bool(re.search(r"\bY\s*=\s*true\b", blk_inner))
            has_b = bool(re.search(r"\bB\s*=\s*true\b", blk_inner))
            sm = re.search(r"Stats\s*=\s*\{", blk_inner)
            stats = {}
            if sm:
                st_open = sm.end() - 1
                st_inner, _ = _lua_brace_span(blk_inner, st_open)
                if st_inner is not None:
                    stats = _lua_parse_stats_table(st_inner)
            if not stats:
                continue
            cat = _gem_categorize(has_r, has_y, has_b, is_meta)
            if not cat:
                continue
            label = format_gem_stats_label(stats)
            if not label:
                continue
            dedupe_key = (cat, label)
            if dedupe_key in seen:
                continue
            seen.add(dedupe_key)
            by_cat[cat].append(
                {
                    "id": gid,
                    "category": cat,
                    "stats": stats,
                    "label": _finalize_gem_label_for_sheet(label),
                }
            )

    for c in by_cat:
        by_cat[c].sort(key=lambda x: _gem_sort_key_label(x["label"]))
    return by_cat


def refresh_pawn_scales_json():
    """Rebuild data/pawn_scales_tbc.json from AddonReference/Pawn/ClassicHawsJon.lua (TBC multipliers)."""
    if not os.path.isfile(PAWN_CLASSIC_HAWS_PATH):
        log(
            "  Pawn scales skipped: missing %s (copy from the Pawn addon if needed)."
            % PAWN_CLASSIC_HAWS_PATH
        )
        return False
    try:
        with open(PAWN_CLASSIC_HAWS_PATH, "r", encoding="utf-8", errors="replace") as f:
            lua_text = f.read()
        scales = _pawn_extract_scales_from_lua(lua_text)
        if not scales:
            log("  Pawn scales: could not parse Classic/TBC block in ClassicHawsJon.lua.")
            return False
        payload = {
            "source": (
                "AddonReference/Pawn/ClassicHawsJon.lua "
                "(Classic + Burning Crusade Classic branch, HawsJon)"
            ),
            "rating_multipliers": (
                "TBC (all *Per variables = 1). For Classic Era, use multipliers "
                "from the Lua file."
            ),
            "scales": scales,
        }
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(PAWN_SCALES_JSON_PATH, "w", encoding="utf-8") as out:
            json.dump(payload, out, indent=2)
        log("  Pawn scales -> %s (%d scales)" % (PAWN_SCALES_JSON_PATH, len(scales)))
        return True
    except OSError as e:
        log("  Pawn scales error: %s" % e)
        return False


def _populate_general_info_sheet(
    ws_gi,
    get_column_letter,
    section_fill,
    section_font,
    intro_title_font,
    intro_body_font,
    subsection_font,
    wrap_align,
):
    """
    First-tab documentation for Google Sheets (Apps Script). Single scrollable column group A:F.
    """
    from openpyxl.styles import Alignment

    last_col = 6
    last_letter = get_column_letter(last_col)
    top_align = Alignment(wrap_text=True, vertical="top", horizontal="left")
    section_align = Alignment(wrap_text=True, vertical="center", horizontal="left")

    for ci in range(1, last_col + 1):
        ws_gi.column_dimensions[get_column_letter(ci)].width = 14

    def span(r1):
        return "A%d:%s%d" % (r1, last_letter, r1)

    r = 1

    def row_section(title):
        nonlocal r
        ws_gi.merge_cells(span(r))
        c = ws_gi.cell(row=r, column=1, value=title)
        c.font = section_font
        c.fill = section_fill
        c.alignment = section_align
        r += 1

    def row_subheading(text):
        nonlocal r
        ws_gi.merge_cells(span(r))
        c = ws_gi.cell(row=r, column=1, value=text)
        c.font = subsection_font
        c.alignment = top_align
        r += 1

    def row_body(text):
        nonlocal r
        ws_gi.merge_cells(span(r))
        c = ws_gi.cell(row=r, column=1, value=text)
        c.font = intro_body_font
        c.alignment = top_align
        r += 1

    def row_blank():
        nonlocal r
        r += 1

    ws_gi.merge_cells(span(r))
    t = ws_gi.cell(row=r, column=1, value="BIS Planner")
    t.font = intro_title_font
    t.alignment = Alignment(wrap_text=False, vertical="center", horizontal="left")
    r += 1
    row_blank()

    row_body(
        "NOTE: This relies heavily on an Apps Script for calculations and will not work outside of Google Sheets."
    )
    row_body(
        "NOTE: There are a lot of background calculations going on. It may take around 5–7 seconds "
        "for the comparison percentages to update."
    )
    row_blank()

    row_section("Current Equipment Tab")

    row_subheading("General")
    row_body(
        "This tab stores all items currently equipped for each spec."
    )
    row_body(
        'To equip an item, select a cell in the "Item Name" column and begin typing. A dropdown list '
        "will appear with items of the same type."
    )
    row_body(
        "Select an item and the stats will autofill and are not editable."
    )
    row_body(
        "If gear is equipped here that is in the BIS Planner tab, it will automatically show as equipped "
        "in that tab and comparison values will automatically update for items of the same spec and gear type."
    )
    row_blank()

    row_subheading("Sockets")
    row_body(
        'Items with gem sockets will automatically assume the "best" (highest contribution to gear score) '
        "gems are applied for the spec."
    )
    row_body(
        "This makes an ideal gear score that is useful when comparing to other items with or without sockets."
    )
    row_body(
        "The assumed gem stats can be found in columns C, D, and E (collapsible). These cells will remain "
        "grey until an item with at least one socket is equipped."
    )
    row_blank()

    row_subheading("Custom Items")
    row_body(
        "Items in the BIS Planner database are limited to Loon's pool. If an item equipped is not on the "
        "list, type a custom name and the stats will be editable for that row."
    )
    row_body(
        "Custom items will have the option to define how many normal and meta sockets they have, and gear "
        "score will assume best gems are applied."
    )
    row_blank()

    row_subheading("Weights")
    row_body(
        "Included for each spec is the default stat weights from Pawn, found in the last row of each spec."
    )
    row_body(
        'The weights are editable. If a change is made, the title "Pawn Defaults" will change to Custom. '
        "The default title can be reselected to reload the default weights."
    )
    row_body(
        "Gear Score is used to compare items, and the % Upgrade (shown in the BIS Planner tab) should match "
        "Pawn. It equals the sum of all stats multiplied by their weights plus socket stats."
    )
    row_blank()
    row_blank()

    row_section("BIS Planner Tab")

    row_subheading("General")
    row_body(
        "This tab is a filterable list of all Loon pre-BIS/BIS items for this class."
    )
    row_body(
        "The only portions of this sheet intended to be edited are column A (Interest) and column M "
        "(notes, if desired)."
    )
    row_body(
        "By default, the notes are mainly used to define which dungeon boss drops a given item, or which "
        "faction or honor level is required."
    )
    row_body(
        "For each spec, only one piece of gear may be equipped at a time per slot, except two at a time "
        "for rings and trinkets."
    )
    row_blank()

    row_subheading("Interest")
    row_body(
        "The Interest column can be used to equip items or classify them by interest level (useful for filtering)."
    )
    row_body(
        'If "Equipped" is selected, the corresponding row in the Current Equipment tab will automatically '
        "update, along with all comparisons for the same spec and gear type."
    )
    row_body(
        'If "----" is selected, the cell will shortly become blank.'
    )
    row_blank()

    row_subheading("Filtering")
    row_body(
        "Any column may be filtered and/or sorted using native Google Sheets tools. Multiple filters can be "
        "applied, but only one sort order at a time."
    )
    row_body(
        'To filter or sort, click the upside-down triangle to the right of a column title; a dropdown menu '
        "will appear."
    )
    row_body(
        "Filter by condition allows text-based filters (contains, does not contain, etc.)."
    )
    row_body(
        "Filter by values lists all distinct entries in the column; uncheck values to hide them."
    )
    row_body(
        '"Clear" unchecks/hides all values in the column (so you can then check only what you want to show). '
        '"Select all" checks/shows all values (effectively removing the filter).'
    )
    row_body(
        "Example: Filter Spec to only desired specs; filter Acquisition Type to dungeon drops or quests; "
        "filter Difficulty to Heroic or Normal (leave blank checked to include non-dungeon items); or filter "
        "Dungeons to specific dungeons."
    )
    row_body(
        "Troubleshoot: If no values are shown, click the Filter button twice to reset the filters (located on "
        "the right side of the Google Sheets toolbar)."
    )
    row_blank()

    row_subheading("Comparison columns (Stats, Stat Comparison, and % Upgrade)")
    row_body(
        "The Stats column shows the stats and special attributes of the item for that row."
    )
    row_body(
        "The Stat Comparison column shows the stat differential between the row item and the currently "
        "equipped item. Positive means the row item's stat is larger."
    )
    row_body(
        "For rings and trinkets, differentials and special attributes are shown for both equipped slots."
    )
    row_body(
        "% Upgrade compares gear scores using: 100 × (item gear score − equipped gear score) ÷ equipped gear score. "
        "For rings and trinkets, both slots are shown on separate lines in the % Upgrade column."
    )
    row_body(
        "If an equipped item with no gear score (i.e. trinket) is compared with an item that has a gear score, "
        "the default % Upgrade is 100%."
    )


def _get_row_color(row):
    acq = row["Acquisition Type"]
    dungeon = row.get("Dungeon", "")
    difficulty = row.get("Difficulty", "")

    if acq == "Dungeon Drop" and dungeon:
        dname = dungeon.replace("Mana-Tombs", "Mana Tombs")
        is_heroic = difficulty == "Heroic"
        palette = DUNGEON_COLORS_HEROIC if is_heroic else DUNGEON_COLORS_NORMAL
        color = palette.get(dname, "D89838" if is_heroic else "FFF0CC")
        return color, is_heroic

    return ACQ_COLORS.get(acq), False


def export_xlsx(csv_path, spec_order, class_title=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        log("WARNING: openpyxl not installed. Skipping Excel export.")
        log("  Install with: pip install openpyxl")
        return

    if not class_title:
        base = os.path.basename(csv_path)
        m = re.match(r"^(.+)_BIS_Planner\.csv$", base, re.I)
        class_title = m.group(1).strip().title() if m else ""

    refresh_pawn_scales_json()

    item_db = load_item_database()

    all_rows = []
    with open(csv_path, "r") as f:
        for row in csv.DictReader(f):
            r = dict(row)
            legacy_sp = (r.pop("Special", None) or "").strip()
            if legacy_sp:
                st = (r.get("Stats") or "").strip()
                merged = (st + " " + legacy_sp).strip() if st else legacy_sp
                r["Stats"] = _break_use_equip_newlines(merged)
            all_rows.append(r)

    if not all_rows:
        log("No rows to export.")
        return

    csv_specs = {(row.get("Spec") or "").strip() for row in all_rows if (row.get("Spec") or "").strip()}
    unknown = csv_specs - set(spec_order)
    if unknown:
        log(
            "WARNING: CSV Spec column has values not in spec_order (Equip formulas use spec_order): "
            + ", ".join(sorted(unknown))
        )
    missing_in_csv = set(spec_order) - csv_specs
    if missing_in_csv:
        log(
            "WARNING: These spec_order entries have no CSV rows (CE sections still exist): "
            + ", ".join(sorted(missing_in_csv))
        )

    header_fill = PatternFill("solid", fgColor="37474F")
    header_font = Font(bold=True, size=11, color="FFFFFF")
    section_fill = PatternFill("solid", fgColor="546E7A")
    section_font = Font(bold=True, size=12, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="BDBDBD"),
        right=Side(style="thin", color="BDBDBD"),
        top=Side(style="thin", color="BDBDBD"),
        bottom=Side(style="thin", color="BDBDBD"),
    )
    wrap_align = Alignment(wrap_text=True, vertical="top")
    intro_title_font = Font(bold=True, size=14, color="212121")
    intro_body_font = Font(size=10, color="424242")
    subsection_font = Font(bold=True, size=11, color="37474F")

    wb = Workbook()
    wb.remove(wb.active)

    # -- General Info (first tab; documentation only) --
    ws_gi = wb.create_sheet(title="General Info", index=0)
    ws_gi.sheet_properties.tabColor = "004D40"
    _populate_general_info_sheet(
        ws_gi,
        get_column_letter,
        section_fill,
        section_font,
        intro_title_font,
        intro_body_font,
        subsection_font,
        wrap_align,
    )

    # -- ItemDB sheet (hidden) --
    ws_db = wb.create_sheet(title="ItemDB")
    db_headers = (
        ["Name", "Slot", "SockR", "SockY", "SockB", "SockM"]
        + STAT_COLUMNS
        + ["Attributes"]
    )
    db_num_cols = len(db_headers)
    for ci, h in enumerate(db_headers, 1):
        ws_db.cell(row=1, column=ci, value=h)

    gem_by_cat = parse_pawn_gems_burning_crusade()
    gem_categories = ["red", "yellow", "blue", "orange", "purple", "green", "meta"]

    items_by_slot = {}
    db_row = 2
    for str_id, entry in sorted(item_db.items(), key=lambda x: x[1].get("name", "")):
        name = entry.get("name", "")
        slot = entry.get("slot", "")
        stats = entry.get("stats", {})
        special = strip_equip_redundant_with_display_stats(
            stats, entry.get("special", "") or ""
        )
        socks = entry.get("sockets") or {}

        ws_db.cell(row=db_row, column=1, value=name)
        ws_db.cell(row=db_row, column=2, value=slot)
        ws_db.cell(row=db_row, column=3, value=int(socks.get("red") or 0))
        ws_db.cell(row=db_row, column=4, value=int(socks.get("yellow") or 0))
        ws_db.cell(row=db_row, column=5, value=int(socks.get("blue") or 0))
        ws_db.cell(row=db_row, column=6, value=int(socks.get("meta") or 0))
        for si, stat_key in enumerate(STAT_COLUMNS):
            val = stats.get(stat_key)
            if val:
                ws_db.cell(row=db_row, column=ITEMDB_FIRST_STAT_COL + si, value=val)
        ws_db.cell(
            row=db_row,
            column=ITEMDB_FIRST_STAT_COL + len(STAT_COLUMNS),
            value=special,
        )

        items_by_slot.setdefault(slot, []).append(name)
        db_row += 1

    # Per-slot item lists (for dropdown data validation references)
    slot_list_cols = {}
    slot_col_start = db_num_cols + 2
    for si, gear_type in enumerate(GEAR_ORDER):
        col = slot_col_start + si
        ws_db.cell(row=1, column=col, value=gear_type)
        slot_items = sorted(items_by_slot.get(gear_type, []))
        for ri, item_name in enumerate(slot_items):
            ws_db.cell(row=2 + ri, column=col, value=item_name)
        slot_list_cols[gear_type] = (get_column_letter(col), len(slot_items))

    gem_list_cols = {}
    next_col = slot_col_start + len(GEAR_ORDER)
    for cat in gem_categories:
        col = next_col
        ws_db.cell(row=1, column=col, value="Gem_%s_label" % cat)
        ws_db.cell(row=1, column=col + 1, value="Gem_%s_id" % cat)
        lst = list(gem_by_cat.get(cat, []))
        # Sentinel only in label column (no numeric id — avoids "0" leaking into dropdowns).
        rows_out = [{"label": _finalize_gem_label_for_sheet("----"), "id": ""}] + lst
        for ri, ent in enumerate(rows_out):
            ws_db.cell(row=2 + ri, column=col, value=ent["label"])
            ws_db.cell(row=2 + ri, column=col + 1, value=ent["id"] if ent["id"] != "" else None)
        ltr = get_column_letter(col)
        ltr_id = get_column_letter(col + 1)
        gem_list_cols[cat] = (ltr, len(rows_out), ltr_id)
        next_col += 2

    ws_db.sheet_state = "hidden"

    # Hidden Id → stats JSON for Apps Script gear score (avoids parsing dropdown labels).
    ws_gems = wb.create_sheet(title="GemDB")
    ws_gems.append(["Id", "Category", "Label", "StatsJson"])
    g_row = 2
    for cat in gem_categories:
        for ent in gem_by_cat.get(cat, []):
            ws_gems.cell(row=g_row, column=1, value=ent["id"])
            ws_gems.cell(row=g_row, column=2, value=cat)
            ws_gems.cell(row=g_row, column=3, value=ent["label"])
            ws_gems.cell(
                row=g_row,
                column=4,
                value=json.dumps(ent["stats"], separators=(",", ":")),
            )
            g_row += 1
    ws_gems.sheet_state = "hidden"
    total_db_items = db_row - 2
    log(f"  ItemDB: {total_db_items} items")

    db_end_col_letter = get_column_letter(db_num_cols)
    itemdb_last_row = max(2, db_row - 1)
    itemdb_range_bounded = f"ItemDB!A$2:{db_end_col_letter}${itemdb_last_row}"

    # -- Current Equipment sheet --
    ws_ce = wb.create_sheet(title="Current Equipment")
    ws_ce.sheet_properties.tabColor = "455A64"

    ce_headers = ["Gear Type", "Item Name"] + CE_GEM_HEADERS + STAT_COLUMNS + ["Gear Score"]

    # Row 1 = column headers (matches apps_script.gs CE_FIRST_DATA_ROW = 2 for first content row)
    CE_HEADER_ROW = 1

    for ci, h in enumerate(ce_headers, 1):
        cell = ws_ce.cell(row=CE_HEADER_ROW, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    ws_ce.column_dimensions["A"].width = 14
    ws_ce.column_dimensions["B"].width = 34
    for ci in range(CE_GEM_FIRST_COL, CE_GEM_LAST_COL + 1):
        letter = get_column_letter(ci)
        ws_ce.column_dimensions[letter].width = 16
    for ci in range(CE_STAT_FIRST_COL, CE_STAT_LAST_COL + 1):
        ws_ce.column_dimensions[get_column_letter(ci)].width = 9
    ws_ce.column_dimensions[get_column_letter(CE_GEAR_SCORE_COL)].width = 11
    for ci in range(CE_PAWN_SNAPSHOT_FIRST_COL, CE_META_WEIGHT_COL + 1):
        letter = get_column_letter(ci)
        ws_ce.column_dimensions[letter].hidden = True
        ws_ce.column_dimensions[letter].width = 2

    # Freeze only Gear Type + Item Name (A:B). Gem block scrolls horizontally with stats (avoids Sheets viewport errors).
    ws_ce.freeze_panes = "%s%d" % (
        get_column_letter(CE_GEM_FIRST_COL),
        CE_HEADER_ROW + 1,
    )

    weight_mode_dv = DataValidation(
        type="list",
        formula1='"%s,%s"' % (CE_WEIGHT_MODE_PAWN, CE_WEIGHT_MODE_CUSTOM),
        allow_blank=False,
    )
    ws_ce.add_data_validation(weight_mode_dv)

    ce_spec_rows = {}
    ce_row = CE_HEADER_ROW + 1
    for spec_name in spec_order:
        cell = ws_ce.cell(row=ce_row, column=1, value=f"--- {spec_name.upper()} ---")
        cell.font = section_font
        cell.fill = section_fill
        for ci in range(2, CE_GEAR_SCORE_COL + 1):
            c = ws_ce.cell(row=ce_row, column=ci)
            c.fill = section_fill
        ce_row += 1

        spec_start = ce_row
        _wd, wvec = get_pawn_weights_for_spec(class_title, spec_name)
        if wvec is None:
            wvec = [0.0] * len(STAT_COLUMNS)
        meta_w = float((_wd or {}).get("MetaSocketEffect") or 0)

        for gear_type in CE_GEAR_ORDER:
            ws_ce.cell(row=ce_row, column=1, value=gear_type)

            for hi in range(len(CE_GEM_HEADERS)):
                ci = CE_GEM_FIRST_COL + hi
                ws_ce.cell(row=ce_row, column=ci, value="")

            for si, stat_key in enumerate(STAT_COLUMNS):
                db_col = ITEMDB_FIRST_STAT_COL + si
                formula = (
                    f'=IFERROR(VLOOKUP(B{ce_row},{itemdb_range_bounded},'
                    f'{db_col},FALSE),"")'
                )
                ws_ce.cell(
                    row=ce_row,
                    column=CE_STAT_FIRST_COL + si,
                    value=formula,
                )

            ws_ce.cell(row=ce_row, column=CE_GEAR_SCORE_COL, value="")

            itemdb_slot = CE_ROW_TO_ITEMDB_SLOT.get(gear_type, gear_type)
            col_info = slot_list_cols.get(itemdb_slot)
            if col_info:
                sl_col, sl_count = col_info
                if sl_count > 0:
                    dv = DataValidation(
                        type="list",
                        formula1=f"ItemDB!${sl_col}$2:${sl_col}${sl_count + 1}",
                        allow_blank=True,
                    )
                    dv.prompt = f"Select or type a {itemdb_slot} item"
                    dv.promptTitle = itemdb_slot
                    ws_ce.add_data_validation(dv)
                    dv.add(ws_ce.cell(row=ce_row, column=2))

            for ci in range(1, CE_GEAR_SCORE_COL + 1):
                ws_ce.cell(row=ce_row, column=ci).border = thin_border

            ce_row += 1

        # Weights row (per spec): B = mode; stats = weights; hidden snapshot + stash + meta weight.
        ws_ce.cell(row=ce_row, column=1, value=CE_WEIGHT_ROW_LABEL)
        ws_ce.cell(row=ce_row, column=2, value=CE_WEIGHT_MODE_PAWN)
        weight_mode_dv.add(ws_ce.cell(row=ce_row, column=2))
        for ci in range(CE_GEM_FIRST_COL, CE_GEM_LAST_COL + 1):
            ws_ce.cell(row=ce_row, column=ci, value="")

        for si, wt in enumerate(wvec):
            ws_ce.cell(
                row=ce_row,
                column=CE_STAT_FIRST_COL + si,
                value=wt,
            )
        ws_ce.cell(row=ce_row, column=CE_GEAR_SCORE_COL, value="")

        for si, wt in enumerate(wvec):
            ws_ce.cell(
                row=ce_row,
                column=CE_PAWN_SNAPSHOT_FIRST_COL + si,
                value=wt,
            )
            ws_ce.cell(
                row=ce_row,
                column=CE_CUSTOM_STASH_FIRST_COL + si,
                value=wt,
            )
        ws_ce.cell(row=ce_row, column=CE_META_WEIGHT_COL, value=meta_w)

        for ci in range(1, CE_META_WEIGHT_COL + 1):
            ws_ce.cell(row=ce_row, column=ci).border = thin_border

        ce_row += 1

        ce_spec_rows[spec_name] = (spec_start, ce_row - 1)
        ce_row += 1

    credit_font = Font(size=8, italic=True, color="999999")
    credit_cell = ws_ce.cell(row=ce_row + 1, column=1, value="Created by Steven Bennett 2026")
    credit_cell.font = credit_font

    log(
        "  Current Equipment: %d specs x (%d gear rows + Weights per spec)"
        % (len(spec_order), len(CE_GEAR_ORDER))
    )

    # -- Single BIS Planner sheet --
    # A–M: through Notes. Equip / CmpRaw / planner gear scores are resolved in Apps Script only (no N–R columns).
    ws = wb.create_sheet(title="BIS Planner")
    ws.sheet_properties.tabColor = "1565C0"

    bp_total_cols = len(SHEET_FIELDS)
    bp_visible_last_col = get_column_letter(SHEET_FIELDS.index("Notes") + 1)

    # A/B: ~header width + small margin for the Sheets filter (Excel char units).
    # Column L: ~"% Upgrade" text width plus space for the filter icon.
    col_widths = {
        "A": 13,
        "B": 10,
        "C": 13,
        "D": 28,
        "E": 10,
        "F": 20,
        "G": 16,
        "H": 13,
        "I": 12,
        "J": 34,
        "K": 34,
        "L": 13.5,
        "M": 14,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.column_dimensions["E"].width = 16

    # Row 1 = column headers (matches apps_script.gs BIS_FIRST_DATA_ROW = 2 for first item row)
    BP_HEADER_ROW = 1

    for ci, field in enumerate(SHEET_FIELDS, 1):
        cell = ws.cell(row=BP_HEADER_ROW, column=ci, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = wrap_align

    ws.freeze_panes = f"E{BP_HEADER_ROW + 1}"
    filter_end = get_column_letter(bp_total_cols)

    interest_dv_std = DataValidation(
        type="list",
        formula1=f'"{",".join(INTEREST_OPTIONS_STANDARD)}"',
        allow_blank=True,
    )
    interest_dv_rt = DataValidation(
        type="list",
        formula1=f'"{",".join(INTEREST_OPTIONS_RING_TRINKET)}"',
        allow_blank=True,
    )
    ws.add_data_validation(interest_dv_std)
    ws.add_data_validation(interest_dv_rt)

    for i, data_row in enumerate(all_rows):
        excel_row = i + BP_HEADER_ROW + 1
        color_hex, dark_bg = _get_row_color(data_row)
        fill = PatternFill("solid", fgColor=color_hex) if color_hex else None
        font = Font(color="FFFFFF", size=10) if dark_bg else Font(size=10)

        for ci, field in enumerate(SHEET_FIELDS, 1):
            if field == "Stat Comparison":
                # Google Sheets: Apps Script writes rich text to K from CE + ItemDB (no CmpRaw column).
                cell = ws.cell(row=excel_row, column=ci, value="")
            elif field == "% Upgrade":
                cell = ws.cell(row=excel_row, column=ci, value="")
            elif field == "Interest":
                cell = ws.cell(row=excel_row, column=ci, value="")
                gear_type = data_row.get("Gear Type", "")
                if gear_type in ("Ring", "Trinket"):
                    interest_dv_rt.add(cell)
                else:
                    interest_dv_std.add(cell)
            else:
                cell = ws.cell(row=excel_row, column=ci, value=data_row.get(field, ""))
            cell.border = thin_border
            cell.alignment = wrap_align
            cell.font = font
            if fill:
                cell.fill = fill

    bp_last_row = BP_HEADER_ROW + len(all_rows)
    ws.auto_filter.ref = f"A{BP_HEADER_ROW}:{filter_end}{bp_last_row}"

    log(f"  BIS Planner: {len(all_rows)} items")

    # Tab order: General Info, Current Equipment, BIS Planner, ItemDB + GemDB (hidden).
    target_order = [
        "General Info",
        "Current Equipment",
        "BIS Planner",
        "ItemDB",
        "GemDB",
    ]
    for idx, name in enumerate(target_order):
        if name in wb.sheetnames:
            current_idx = wb.sheetnames.index(name)
            wb.move_sheet(name, offset=idx - current_idx)

    xlsx_path = csv_path.rsplit(".", 1)[0] + ".xlsx"
    wb.save(xlsx_path)
    log(f"Excel -> {xlsx_path} ({len(all_rows)} items, {total_db_items} in ItemDB)")


# ============================================================
# MAIN
# ============================================================


def run_stat_audit_report():
    """
    Data audit: data/item_database.json stat keys (frequency) vs STAT_COLUMNS and Pawn scale keys.
    Run: python3 generate_bis_planner.py --audit-stats
    """
    stat_cols_set = set(STAT_COLUMNS)
    pawn_mapped = set()
    for tup in STAT_COL_PAWN_WEIGHT_KEYS:
        pawn_mapped.update(tup)

    if not os.path.isfile(ITEM_DB_PATH):
        log("ERROR: %s not found." % ITEM_DB_PATH)
        return
    with open(ITEM_DB_PATH, "r", encoding="utf-8") as f:
        item_db = json.load(f)
    key_freq = Counter()
    for entry in item_db.values():
        for k in (entry.get("stats") or {}):
            key_freq[k] += 1

    log("--- ItemDB stats: top keys by item count ---")
    for k, n in key_freq.most_common(50):
        in_col = "CE column" if k in stat_cols_set else "not in STAT_COLUMNS"
        log("  %6d  %-28s  %s" % (n, k, in_col))

    orphan_stats = sorted(k for k in key_freq if k not in stat_cols_set and not k.startswith("_"))
    if orphan_stats:
        log("--- ItemDB stat keys not mapped to STAT_COLUMNS (excluding _internal) ---")
        for k in orphan_stats:
            log("  %6d  %s" % (key_freq[k], k))

    pawn_keys_max = {}
    if os.path.isfile(PAWN_SCALES_JSON_PATH):
        with open(PAWN_SCALES_JSON_PATH, "r", encoding="utf-8") as f:
            payload = json.load(f)
        for sc in payload.get("scales", []):
            w = sc.get("weights") or {}
            for pk, val in w.items():
                v = abs(float(val or 0))
                if v > pawn_keys_max.get(pk, 0):
                    pawn_keys_max[pk] = v
    unmapped_pawn = sorted(
        pk for pk, mx in pawn_keys_max.items() if mx > 0 and pk not in pawn_mapped
    )
    if unmapped_pawn:
        log("--- Pawn keys with non-zero weight somewhere, not aggregated into CE stat columns ---")
        for pk in unmapped_pawn:
            log("  max|w|≈%.4g  %s" % (pawn_keys_max[pk], pk))
    log("--- Audit done (STAT_COLUMNS has %d entries) ---" % len(STAT_COLUMNS))


def run_loon_pawn_spec_check():
    """Print Loon RegisterSpec labels vs data/pawn_scales_tbc.json spec_name (per class)."""
    if not os.path.isdir(LOON_GUIDES_DIR):
        log("ERROR: Loon dir missing: %s" % LOON_GUIDES_DIR)
        return
    if not os.path.isfile(PAWN_SCALES_JSON_PATH):
        log("ERROR: %s not found." % PAWN_SCALES_JSON_PATH)
        return
    with open(PAWN_SCALES_JSON_PATH, "r", encoding="utf-8") as f:
        payload = json.load(f)
    pawn_by_class = {}
    for sc in payload.get("scales", []):
        cn = (sc.get("class_name") or "").strip()
        sn = (sc.get("spec_name") or "").strip()
        pawn_by_class.setdefault(cn, set()).add(sn)

    loon_by_class = {}
    loon_resolved_by_class = {}
    loon_files = []
    for fname in sorted(os.listdir(LOON_GUIDES_DIR)):
        if not fname.endswith(".lua"):
            continue
        path = os.path.join(LOON_GUIDES_DIR, fname)
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for line in f:
                m = re.match(
                    r'.*LBIS:RegisterSpec\(LBIS\.L\["([^"]*)"\],\s*LBIS\.L\["([^"]*)"\]',
                    line,
                )
                if m:
                    c, s = m.group(1), m.group(2)
                    loon_by_class.setdefault(c, set()).add(s)
                    cls_title = c.strip().title()
                    s_clean = (s or "").strip()
                    resolved = PLANNER_SPEC_TO_PAWN_SPEC_NAME.get((cls_title, s_clean), s_clean)
                    loon_resolved_by_class.setdefault(c, set()).add(resolved)
                    loon_files.append((fname, c, s))
                    break

    log("--- Loon -> Pawn: mismatches after PLANNER_SPEC_TO_PAWN_SPEC_NAME ---")
    bad = 0
    for fname, c, s in sorted(loon_files, key=lambda x: (x[1], x[2], x[0])):
        cls_title = c.strip().title()
        s_clean = (s or "").strip()
        want_pawn = PLANNER_SPEC_TO_PAWN_SPEC_NAME.get((cls_title, s_clean), s_clean)
        pset = pawn_by_class.get(c, set())
        if want_pawn not in pset:
            bad += 1
            log(
                "  %s  class=%r loon=%r -> %r  Pawn has: %s"
                % (fname, c, s_clean, want_pawn, sorted(pset))
            )
    if bad == 0:
        log("  (none — all Loon specs resolve to a Pawn scale)")

    log("--- Pawn scales with no Loon guide (same class; Loon labels resolved via alias map) ---")
    for cn in sorted(pawn_by_class.keys()):
        only_pawn = pawn_by_class[cn] - loon_resolved_by_class.get(cn, set())
        if only_pawn:
            log("  %s: %s" % (cn, ", ".join(sorted(only_pawn))))

    log("--- Done ---")


def list_guide_classes():
    """Lowercase class names that have at least one guide file in AddonReference/Loon/."""
    classes = set()
    for fname in os.listdir(LOON_GUIDES_DIR):
        if fname.endswith(".lua"):
            m = re.match(r"([A-Z][a-z]+)", fname)
            if m:
                classes.add(m.group(1).lower())
    return sorted(classes)


def main():
    parser = argparse.ArgumentParser(
        description="BIS Planner — TBC Classic Gear Spreadsheet Generator",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Workflow:\n"
            "  1. python3 generate_bis_planner.py --build-db       # fetch item stats\n"
            "     (auto re-fetches cache rows with empty stats+special; use --refresh-db for full refetch)\n"
            "  2. python3 generate_bis_planner.py paladin           # generate Paladin sheet\n"
            "     python3 generate_bis_planner.py --all              # generate every class\n"
            "     python3 generate_bis_planner.py all               # same as --all\n"
            "\n"
            "  Classes: druid, hunter, mage, paladin, priest, rogue, shaman, warlock, warrior\n"
            "  Phase 0 (default) = PreRaid. Phase 1 = PreRaid + Phase 1 items (cumulative).\n"
        ),
    )
    parser.add_argument("class_name", nargs="?", default=None,
                        help="Class to generate (e.g., paladin, warrior, druid), or 'all' for every class")
    parser.add_argument("--all", action="store_true",
                        help="Generate CSV and Excel for every class that has guide files")
    parser.add_argument("--build-db", action="store_true",
                        help="Fetch item stats from Wowhead and build data/item_database.json")
    parser.add_argument("--refresh-db", action="store_true",
                        help="With --build-db: refetch every catalog item (ignore cache hits)")
    parser.add_argument("--test", action="store_true",
                        help="With --build-db: only fetch items for the specified class (or all guide items)")
    parser.add_argument("--phase", type=int, default=0,
                        help="Max phase to include (cumulative). 0=PreRaid (default), 1=Phase 1, etc.")
    parser.add_argument("--output", type=str, default=None, help="Output CSV path override")
    parser.add_argument("--no-excel", action="store_true", help="Skip .xlsx generation")
    parser.add_argument(
        "--audit-stats",
        action="store_true",
        help="Print data/item_database.json stat-key frequencies vs STAT_COLUMNS / Pawn keys; then exit",
    )
    parser.add_argument(
        "--check-loon-pawn-specs",
        action="store_true",
        help="Compare Loon RegisterSpec names to Pawn JSON spec_name per class; then exit",
    )
    args = parser.parse_args()

    if args.audit_stats:
        run_stat_audit_report()
        return

    if args.check_loon_pawn_specs:
        run_loon_pawn_spec_check()
        return

    if args.build_db:
        build_item_database(
            test_mode=args.test,
            test_class=args.class_name,
            force_refresh=args.refresh_db,
        )
        return

    run_all = args.all or (
        args.class_name is not None and args.class_name.strip().lower() == "all"
    )
    if run_all:
        if args.output:
            log("ERROR: --output cannot be used with --all or 'all' (each class writes its own file).")
            sys.exit(1)
        classes = list_guide_classes()
        log(f"Generating all {len(classes)} classes: {', '.join(classes)}")
        for cn in classes:
            class_title = cn.strip().title()
            output_csv = os.path.join(OUTPUT_DIR, f"{class_title}_BIS_Planner.csv")
            rows, spec_order = generate_csv(args.phase, cn, output_csv)
            if not args.no_excel:
                export_xlsx(output_csv, spec_order, class_title=class_title)
        log("Done (all classes).")
        return

    if not args.class_name:
        available = list_guide_classes()
        log("ERROR: Please specify a class name.")
        log(f"  Available classes: {', '.join(available)}")
        log(f"  Example: python3 generate_bis_planner.py paladin")
        log(f"  Or all: python3 generate_bis_planner.py --all")
        sys.exit(1)

    class_title = args.class_name.strip().title()
    output_csv = args.output or os.path.join(OUTPUT_DIR, f"{class_title}_BIS_Planner.csv")

    rows, spec_order = generate_csv(args.phase, args.class_name, output_csv)

    if not args.no_excel:
        export_xlsx(output_csv, spec_order, class_title=class_title)


if __name__ == "__main__":
    main()
